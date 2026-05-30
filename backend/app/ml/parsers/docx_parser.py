import structlog

logger = structlog.get_logger(__name__)


def parse_docx(file_path: str) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        from docx import Document  # python-docx
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX parsing. Install it with: pip install python-docx") from exc

    doc = Document(file_path)
    parts: list[str] = []

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Extract table cells
    for table in doc.tables:
        for row in table.rows:
            row_texts: list[str] = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                parts.append(" | ".join(row_texts))

    result = "\n".join(parts)
    logger.info("docx_parsed", file_path=file_path, paragraphs=len(doc.paragraphs), chars=len(result))
    return result
